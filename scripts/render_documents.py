"""Create editable DOCX and PDF using local Pandoc/LibreOffice."""
import hashlib
import json
import os
import sys
from pathlib import Path
import subprocess
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
items=[('manuscript/main.md','paper1-manuscript'),('manuscript/supplement.md','paper1-supplement'),
       ('dossier/cover-letter.md','paper1-cover-letter'),('dossier/skeptical-review-packet.md','paper1-skeptical-review')]
receipt=[]
for source,name in items:
    p=ROOT/source; before=hashlib.sha256(p.read_bytes()).hexdigest()
    args=['pandoc',str(p),'-f','markdown+tex_math_single_backslash+tex_math_dollars',
          '--resource-path',str(ROOT/'manuscript')+':'+str(ROOT),
          '-o',str(OUT/f'{name}.docx')]
    subprocess.run(args,cwd=ROOT,check=True)
    document=Document(OUT/f'{name}.docx')
    for paragraph in document.paragraphs:
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_with_next=True
        if paragraph.style.name=='Image Caption':
            paragraph.paragraph_format.keep_together=True
            for run in paragraph.runs:
                run.italic=False
    for section in document.sections:
        footer=section.footer.paragraphs[0]
        footer.alignment=2
        run=footer.add_run('Internal review • ');run.font.size=Pt(8)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for table in document.tables:
        headers=[cell.text for cell in table.rows[0].cells]
        if headers[0]=='Stage': weights=[3.5,.75,1.,1.25]
        elif len(headers)==2: weights=[4.5,2.]
        elif len(headers)==3: weights=[3.,1.75,1.75]
        elif headers[0]=='Dataset' and len(headers)==5: weights=[.9,2.1,1.2,1.15,1.15]
        elif headers[0]=='Dataset and reference': weights=[1.7,1.2,1.8,1.8]
        elif len(headers)==4: weights=[2.8,1.3,1.1,1.3]
        elif len(headers)==5: weights=[2.5,1.,1.,1.,1.]
        else: weights=[6.5/len(headers)]*len(headers)
        table.autofit=False
        for column,width in zip(table.columns,weights):column.width=Inches(width)
        for row in table.rows:
            props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,weights):
                cell.width=Inches(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(3)
                    for run in paragraph.runs:run.font.size=Pt(9)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    document.save(OUT/f'{name}.docx')
    assert hashlib.sha256(p.read_bytes()).hexdigest()==before,'Source changed during render'
    receipt.append({'source':source,'source_sha256':before,'docx':name+'.docx'})
profile=ROOT/'work/libreoffice-review-profile'
environment=os.environ.copy()
if '--local-ipc-workaround' in sys.argv:
    ipc=ROOT/'work/doc-ipc';ipc.mkdir(parents=True,exist_ok=True)
    environment['TMPDIR']=str(ipc)
    shim=ROOT/'work/libreoffice_socket_path.so'
    subprocess.run(['gcc','-shared','-fPIC','-Wall','-Wextra','-o',str(shim),
                    str(ROOT/'scripts/libreoffice_socket_path.c'),'-ldl'],
                   cwd=ROOT,env=environment,check=True)
    environment['LD_PRELOAD']=str(shim)
command=['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(OUT),*[str(OUT/f'{name}.docx') for _,name in items]]
converted=subprocess.run(command,cwd=ROOT,env=environment,capture_output=True,text=True,timeout=180)
(ROOT/'receipts/m4-render.log').write_text(converted.stdout+converted.stderr)
converted.check_returncode()
for r,(_,name) in zip(receipt,items):
    pdf=OUT/f'{name}.pdf'
    if not pdf.exists() or pdf.stat().st_size<1000: raise RuntimeError('PDF conversion missing: '+name)
    subprocess.run(['pdftotext','-layout',str(pdf),str(OUT/f'{name}.txt')],check=True)
    info=subprocess.run(['pdfinfo',str(pdf)],capture_output=True,text=True,check=True).stdout
    r['pdf']=name+'.pdf';r['pdf_info']=info
    r['docx_sha256']=hashlib.sha256((OUT/r['docx']).read_bytes()).hexdigest()
    r['pdf_sha256']=hashlib.sha256(pdf.read_bytes()).hexdigest()
(ROOT/'receipts/m4-render.json').write_text(json.dumps({'created_at':datetime.now(timezone.utc).isoformat(),'local_ipc_workaround':'--local-ipc-workaround' in sys.argv,'documents':receipt},indent=2)+'\n')
print(json.dumps({'documents':len(receipt),'output':str(OUT)}))
